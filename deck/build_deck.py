# -*- coding: utf-8 -*-
"""AI Data Conversion Studio - PwC-themed BUSINESS overview deck (python-pptx)."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn

# ---------- PwC palette ----------
ORANGE     = RGBColor(0xD0, 0x4A, 0x02)
ORANGE_DK  = RGBColor(0xA8, 0x3A, 0x00)
TANGERINE  = RGBColor(0xEB, 0x8C, 0x00)
YELLOW     = RGBColor(0xFF, 0xB6, 0x00)
RED        = RGBColor(0xE0, 0x30, 0x1E)
ROSE       = RGBColor(0xC0, 0x27, 0x1F)
GREEN      = RGBColor(0x17, 0x8A, 0x4C)
TEAL       = RGBColor(0x1C, 0x72, 0x93)
INK        = RGBColor(0x2B, 0x2B, 0x2B)
INK2       = RGBColor(0x1C, 0x1C, 0x1C)
PANEL_DK   = RGBColor(0x2E, 0x2E, 0x2E)
GRAY       = RGBColor(0x5A, 0x5A, 0x5A)
MUTE       = RGBColor(0x8A, 0x8A, 0x8A)
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
SOFT       = RGBColor(0xFD, 0xEE, 0xE4)
SAND       = RGBColor(0xF6, 0xF3, 0xF0)
LINE       = RGBColor(0xE4, 0xDE, 0xD9)
LGRAY      = RGBColor(0xC9, 0xC9, 0xC9)
DGRAY      = RGBColor(0xD6, 0xD6, 0xD6)
BORDER_DK  = RGBColor(0x3D, 0x3D, 0x3D)
SOFTBORDER = RGBColor(0xF3, 0xC9, 0xA6)

HEAD = "Arial"
BODY = "Arial"

EMU = 914400
W, H = 13.333, 7.5

prs = Presentation()
prs.slide_width  = Emu(int(W * EMU))
prs.slide_height = Emu(int(H * EMU))
BLANK = prs.slide_layouts[6]


# ---------- helpers ----------
def slide():
    return prs.slides.add_slide(BLANK)

def bg(s, color):
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = color

def _no_line(shp):
    shp.line.fill.background()

def _line(shp, color, w=1.0):
    shp.line.color.rgb = color
    shp.line.width = Pt(w)

def _empty_effect():
    from pptx.oxml import parse_xml
    return parse_xml('<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>')

def rect(s, x, y, w, h, color, rounded=False, line_color=None, line_w=1.0,
         radius=0.09, rot=0):
    from pptx.enum.shapes import MSO_SHAPE
    shp_type = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = s.shapes.add_shape(shp_type, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if color is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = color
    if line_color is None:
        _no_line(shp)
    else:
        _line(shp, line_color, line_w)
    if rot:
        shp.rotation = rot
    shp._element.spPr.append(_empty_effect())
    return shp

def ellipse(s, x, y, d, color, line_color=None):
    from pptx.enum.shapes import MSO_SHAPE
    shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    if line_color is None:
        _no_line(shp)
    else:
        _line(shp, line_color, 1)
    shp._element.spPr.append(_empty_effect())
    return shp

def arrow(s, x, y, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE
    shp = s.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    _no_line(shp)
    shp._element.spPr.append(_empty_effect())
    return shp

def _set_run(r, text, size, color, bold, face, spacing=None):
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.bold = bold
    r.font.name = face
    if spacing is not None:
        rPr = r._r.get_or_add_rPr()
        rPr.set('spc', str(int(spacing * 100)))

def text(s, x, y, w, h, txt, size=12, color=INK, bold=False, face=BODY,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None,
         line_mult=None, wrap=True):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    para = tf.paragraphs[0]
    para.alignment = align
    if line_mult:
        para.line_spacing = line_mult
    r = para.add_run()
    _set_run(r, txt, size, color, bold, face, spacing)
    return tb

def rich(s, x, y, w, h, segments, size=12, align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, line_mult=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    para = tf.paragraphs[0]
    para.alignment = align
    if line_mult:
        para.line_spacing = line_mult
    for (t, c, b) in segments:
        r = para.add_run()
        _set_run(r, t, size, c, b, BODY)
    return tb

def _add_bullet(para, color):
    from pptx.oxml import parse_xml
    pPr = para._p.get_or_add_pPr()
    pPr.set('marL', str(int(0.22 * EMU)))
    pPr.set('indent', str(int(-0.22 * EMU)))
    ns = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    pPr.append(parse_xml('<a:buClr %s><a:srgbClr val="%02X%02X%02X"/></a:buClr>' % (ns, color[0], color[1], color[2])))
    pPr.append(parse_xml('<a:buFont %s typeface="Arial"/>' % ns))
    pPr.append(parse_xml('<a:buChar %s char="%s"/>' % (ns, "•")))

def bullets(s, x, y, w, h, items, size=11, color=INK, gap=8, line_mult=1.05,
            bullet_color=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    bc = bullet_color or ORANGE
    for i, it in enumerate(items):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.line_spacing = line_mult
        para.space_after = Pt(gap)
        r = para.add_run()
        _set_run(r, it, size, color, False, BODY)
        _add_bullet(para, bc)
    return tb

def brand_mark(s, x, y, scale=1.0):
    d = 0.13 * scale
    gap = 0.055 * scale
    for i, c in enumerate([YELLOW, ORANGE, RED]):
        ellipse(s, x + i * (d + gap), y, d, c)

def footer(s, n, on_dark=False):
    c = RGBColor(0x9A, 0x9A, 0x9A) if on_dark else MUTE
    text(s, 0.55, H - 0.42, 6, 0.3, "AI Data Conversion Studio", size=8.5, color=c)
    text(s, W - 3.7, H - 0.42, 2.5, 0.3, "PwC  ·  Confidential", size=8.5, color=c, align=PP_ALIGN.RIGHT)
    text(s, W - 1.05, H - 0.42, 0.5, 0.3, str(n), size=8.5, color=c, align=PP_ALIGN.RIGHT)

def header(s, kicker, title):
    brand_mark(s, 0.55, 0.5, 1.0)
    text(s, 0.55, 0.72, 11.5, 0.3, kicker.upper(), size=11, color=ORANGE, bold=True, spacing=2)
    text(s, 0.53, 1.0, 12.3, 0.75, title, size=30, color=INK, bold=True, face=HEAD)

def content_slide(kicker, title, n):
    s = slide(); bg(s, WHITE)
    header(s, kicker, title)
    footer(s, n, False)
    return s

def feat_card(s, x, y, w, h, badge, badge_color, heading, body, small=False):
    rect(s, x, y, w, h, WHITE, rounded=True, line_color=LINE, line_w=1, radius=0.06)
    bs = 0.42 if small else 0.5
    rect(s, x + 0.2, y + 0.2, bs, bs, badge_color, rounded=True, radius=0.14)
    if badge:
        text(s, x + 0.2, y + 0.2, bs, bs, badge, size=(13 if small else 15),
             color=WHITE, bold=True, face=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    hx = x + 0.2 + bs + 0.1
    text(s, hx, y + 0.16, w - (bs + 0.5), 0.5, heading, size=(12.5 if small else 13.5),
         color=INK, bold=True, face=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.24, y + (0.72 if small else 0.82), w - 0.48, h - (0.86 if small else 1.0),
         body, size=(9.6 if small else 10.5), color=GRAY, line_mult=1.05)

def grid3x2(s, items, gy=2.0):
    cw, ch, gx, gapX, gapY = 3.86, 1.95, 0.55, 0.32, 0.30
    for i, it in enumerate(items):
        col, row = i % 3, i // 3
        x = gx + col * (cw + gapX)
        y = gy + row * (ch + gapY)
        feat_card(s, x, y, cw, ch, it[0], it[3], it[1], it[2], small=True)

def grid2x2(s, items, sand=False, circle=False, gy=2.0):
    cw, ch, gx, gapX, gapY = 5.9, 1.95, 0.55, 0.55, 0.32
    for i, it in enumerate(items):
        col, row = i % 2, i // 2
        x = gx + col * (cw + gapX)
        y = gy + row * (ch + gapY)
        rect(s, x, y, cw, ch, (SAND if sand else WHITE), rounded=True, line_color=LINE, line_w=1, radius=0.06)
        if circle:
            rect(s, x + 0.22, y + 0.24, 0.5, 0.5, it[2], rounded=True, radius=0.5)
            ellipse(s, x + 0.36, y + 0.38, 0.22, WHITE)
        else:
            rect(s, x + 0.22, y + 0.22, 0.5, 0.5, it[2], rounded=True, radius=0.14)
        text(s, x + 0.85, y + 0.2, cw - 1.05, 0.55, it[0], size=14, color=INK,
             bold=True, face=HEAD, anchor=MSO_ANCHOR.MIDDLE)
        text(s, x + 0.26, y + 0.84, cw - 0.5, ch - 1.0, it[1], size=10.8, color=GRAY, line_mult=1.05)


# =====================================================================
# 1 - TITLE
# =====================================================================
s = slide(); bg(s, INK2)
rect(s, 9.7, 4.4, 4.6, 4.6, ORANGE, rounded=True, radius=0.08, rot=20)
rect(s, 10.9, 5.1, 3.6, 3.6, TANGERINE, rounded=True, radius=0.08, rot=20)
rect(s, 11.7, 5.7, 2.6, 2.6, YELLOW, rounded=True, radius=0.08, rot=20)
brand_mark(s, 0.75, 0.75, 1.5)
text(s, 0.75, 1.15, 10, 0.35, "PwC  ·  DATA MIGRATION ACCELERATOR", size=12, color=YELLOW, bold=True, spacing=2)
text(s, 0.72, 2.2, 9.8, 1.6, "AI Data\nConversion Studio", size=52, color=WHITE, bold=True, face=HEAD, line_mult=0.95)
text(s, 0.75, 4.05, 8.4, 0.95, "Convert legacy data to Guidewire faster, with less effort and lower risk",
     size=17, color=DGRAY, line_mult=1.1)
text(s, 0.75, 6.4, 11, 0.4, "A business overview", size=12.5, color=LGRAY)

# =====================================================================
# 2 - AGENDA
# =====================================================================
s = content_slide("Overview", "What we will cover", 2)
items = [
    ["01", "The challenge", "Why converting legacy data to Guidewire is slow, costly and risky today.", ORANGE],
    ["02", "The solution", "What AI Data Conversion Studio is — in plain terms — and how it helps.", TANGERINE],
    ["03", "How it works", "The guided, end-to-end journey from source data to a validated load.", YELLOW],
    ["04", "The business value", "Faster delivery, lower cost, higher quality and full control.", GREEN],
]
cw, ch, gx, gy, gapX, gapY = 5.9, 1.9, 0.55, 2.0, 0.55, 0.35
for i, it in enumerate(items):
    col, row = i % 2, i // 2
    feat_card(s, gx + col*(cw+gapX), gy + row*(ch+gapY), cw, ch, it[0], it[3], it[1], it[2])

# =====================================================================
# 3 - THE CHALLENGE
# =====================================================================
s = content_slide("The challenge", "Data conversion is the riskiest part of a Guidewire program", 3)
text(s, 0.55, 1.72, 12.2, 0.6,
     "Moving decades of legacy policy, claims and billing data into Guidewire is where most time, cost and risk hide — and it is often what holds up go-live.",
     size=13.5, color=GRAY, line_mult=1.1)
pains = [
    ["Slow and manual", "Analysts hand-map thousands of data fields in spreadsheets — weeks of effort that is hard to review and easy to get wrong."],
    ["Legacy data is a black box", "Source data arrives in many shapes and formats with little documentation on what each field actually means."],
    ["Costly rework", "Transformation logic is built by hand and drifts from the plan — small changes ripple into expensive rework."],
    ["Late surprises at go-live", "Duplicates, missing values and bad codes surface only after loading — forcing fire-drills close to go-live."],
]
cw, ch, gx, gy, gapX, gapY = 5.9, 1.62, 0.55, 2.5, 0.55, 0.3
for i, it in enumerate(pains):
    col, row = i % 2, i // 2
    x = gx + col*(cw+gapX); y = gy + row*(ch+gapY)
    rect(s, x, y, cw, ch, SAND, rounded=True, line_color=LINE, line_w=1, radius=0.06)
    ellipse(s, x + 0.24, y + 0.26, 0.34, RED)
    text(s, x + 0.24, y + 0.26, 0.34, 0.34, "!", size=15, color=WHITE, bold=True, face=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.72, y + 0.22, cw - 0.9, 0.4, it[0], size=13.5, color=INK, bold=True, face=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.26, y + 0.68, cw - 0.5, ch - 0.82, it[1], size=10.5, color=GRAY, line_mult=1.05)

# =====================================================================
# 4 - COST OF THE STATUS QUO (stat callouts)
# =====================================================================
s = content_slide("The challenge", "The cost of doing it the old way", 4)
text(s, 0.55, 1.72, 12.2, 0.55,
     "Every manual step adds cost, delay and risk to the program — long before anyone sees whether the data is right.",
     size=13.5, color=GRAY, line_mult=1.1)
stats = [
    ["Weeks", "of manual effort to map each source system, field by field", ORANGE],
    ["1,000s", "of fields to map, transform and review by hand", TANGERINE],
    ["Go-live", "when hidden data-quality issues typically surface — the worst time", RED],
]
cw, gx, gapX, gy, ch = 3.9, 0.55, 0.44, 2.6, 3.0
for i, st in enumerate(stats):
    x = gx + i*(cw+gapX)
    rect(s, x, gy, cw, ch, SAND, rounded=True, line_color=LINE, line_w=1, radius=0.06)
    text(s, x + 0.3, gy + 0.5, cw - 0.6, 1.0, st[0], size=48, color=st[2], bold=True, face=HEAD)
    text(s, x + 0.3, gy + 1.7, cw - 0.6, 1.1, st[1], size=12.5, color=INK, line_mult=1.12)
rect(s, 0.55, 6.05, 12.23, 0.7, SOFT, rounded=True, line_color=SOFTBORDER, line_w=1, radius=0.12)
rich(s, 0.8, 6.05, 11.8, 0.7,
     [("The bottom line:  ", ORANGE_DK, True),
      ("more effort does not mean more confidence — issues are still found late.", INK, False)],
     size=12.5, anchor=MSO_ANCHOR.MIDDLE)

# =====================================================================
# 5 - THE SOLUTION (introduction)
# =====================================================================
s = content_slide("The solution", "Meet AI Data Conversion Studio", 5)
rich(s, 0.55, 1.85, 6.5, 1.05,
     [("A guided, AI-assisted workspace that turns legacy source data into ", INK, False),
      ("Guidewire-ready data", ORANGE_DK, True),
      (" — faster, with your team firmly in control.", INK, False)],
     size=14, line_mult=1.15)
bl = [
    "Bring in your source data and your target — no matter the format.",
    "AI proposes the field-by-field mappings and transformation rules.",
    "Your experts review, refine and approve — nothing is a black box.",
    "The tool generates the code to move the data into the target.",
    "Data quality is checked against the target before go-live.",
]
bullets(s, 0.6, 3.05, 6.5, 3.3, bl, size=12.5, color=GRAY, gap=9)
# right panel: what it does for you
px, pw = 7.6, 5.2
rect(s, px, 1.85, pw, 4.55, INK2, rounded=True, radius=0.05)
text(s, px + 0.35, 2.1, pw - 0.7, 0.4, "What it does for you", size=14, color=YELLOW, bold=True, face=HEAD)
facts = [
    ["Speed", "First-draft mappings in minutes, not weeks"],
    ["Confidence", "Issues caught early, not at go-live"],
    ["Control", "People approve every mapping the AI suggests"],
    ["Traceability", "A complete record of every decision"],
    ["Consistency", "One repeatable way of working across clients"],
]
fy = 2.62
for f in facts:
    ellipse(s, px + 0.35, fy + 0.06, 0.16, ORANGE)
    text(s, px + 0.62, fy - 0.05, pw - 0.95, 0.3, f[0], size=12.5, color=WHITE, bold=True, face=HEAD)
    text(s, px + 0.62, fy + 0.24, pw - 0.95, 0.42, f[1], size=10, color=LGRAY, line_mult=1.0)
    fy += 0.73

# =====================================================================
# 6 - WHERE IT FITS (conceptual flow)
# =====================================================================
s = content_slide("Context", "Where it fits in your migration", 6)
text(s, 0.55, 1.72, 12.2, 0.55,
     "Your legacy data is prepared, transformed and loaded into Guidewire. The Studio is where the mapping and transformation work gets defined — the hard part.",
     size=12.5, color=GRAY, line_mult=1.1)
laneY, laneH, laneW = 2.7, 3.15, 3.55
lanes = [
    ["Your source data", RED, ["Policy · Claims · Billing", "Customer & reference data", "Databases and files"]],
    ["Prepare & transform", ORANGE, ["Organise the data", "Apply the mappings", "Ready it for the target"]],
    ["Guidewire", INK, ["ClaimCenter · PolicyCenter", "Loaded & reconciled", "Ready for business use"]],
]
xs = [0.55, 0.55 + laneW + 1.05, 0.55 + 2*(laneW + 1.05)]
for i, ln in enumerate(lanes):
    x = xs[i]
    rect(s, x, laneY, laneW, laneH, SAND, rounded=True, line_color=LINE, line_w=1, radius=0.05)
    rect(s, x, laneY, laneW, 0.55, ln[1], rounded=True, radius=0.08)
    text(s, x, laneY, laneW, 0.55, ln[0], size=13, color=WHITE, bold=True, face=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    by = laneY + 0.8
    for b in ln[2]:
        rect(s, x + 0.3, by, laneW - 0.6, 0.6, WHITE, rounded=True, line_color=LINE, line_w=1, radius=0.08)
        text(s, x + 0.35, by, laneW - 0.7, 0.6, b, size=10.5, color=INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        by += 0.74
for i in range(2):
    ax = xs[i] + laneW + 0.17
    arrow(s, ax, laneY + laneH/2 - 0.22, 0.7, 0.44, TANGERINE)
rect(s, 0.55, 6.2, 12.23, 0.72, SOFT, rounded=True, line_color=SOFTBORDER, line_w=1, radius=0.12)
rich(s, 0.8, 6.2, 11.8, 0.72,
     [("AI Data Conversion Studio  ", ORANGE_DK, True),
      ("is where you define and approve the mappings and generate what is needed to move the data.", INK, False)],
     size=11.5, anchor=MSO_ANCHOR.MIDDLE)

# =====================================================================
# 7 - HOW IT WORKS (6 phases)
# =====================================================================
s = slide(); bg(s, INK2)
brand_mark(s, 0.75, 0.7, 1.3)
text(s, 0.75, 1.1, 11, 0.35, "HOW IT WORKS", size=12, color=YELLOW, bold=True, spacing=2)
text(s, 0.73, 1.42, 11.5, 0.7, "One guided journey, six simple steps", size=30, color=WHITE, bold=True, face=HEAD)
phases = [
    ["Set up", "Bring in your source data and define the target", ORANGE],
    ["Understand", "See what your data really contains before mapping", TANGERINE],
    ["Map with AI", "Let AI draft the mappings, then review and approve", YELLOW],
    ["Build", "Generate what is needed to move the data", GREEN],
    ["Validate", "Check data quality against the target", TEAL],
    ["Deliver", "Hand off, report and keep a full audit trail", RED],
]
cw, ch, gx, gy, gapX, gapY = 3.86, 1.75, 0.75, 2.5, 0.32, 0.32
for i, ph in enumerate(phases):
    col, row = i % 3, i // 3
    x = gx + col*(cw+gapX); y = gy + row*(ch+gapY)
    rect(s, x, y, cw, ch, PANEL_DK, rounded=True, line_color=BORDER_DK, line_w=1, radius=0.05)
    ellipse(s, x + 0.26, y + 0.28, 0.46, ph[2])
    num_color = INK if i == 2 else WHITE
    text(s, x + 0.26, y + 0.28, 0.46, 0.46, str(i+1), size=15, color=num_color, bold=True, face=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.85, y + 0.26, cw - 1.05, 0.5, ph[0], size=16, color=WHITE, bold=True, face=HEAD, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.28, y + 0.9, cw - 0.55, ch - 1.05, ph[1], size=10.5, color=LGRAY, line_mult=1.1)
footer(s, 7, True)

# =====================================================================
# 8 - CAPABILITIES: PREPARE & UNDERSTAND
# =====================================================================
s = content_slide("What you can do · Steps 1–2", "Prepare your data and understand it", 8)
grid3x2(s, [
    ["A", "Know Your Data", "Upload your documents and simply ask questions about your data in plain language.", ORANGE],
    ["B", "Bring in any source", "Connect a database or upload files in almost any format — the tool reads the structure for you.", TANGERINE],
    ["C", "Define the target", "Point the tool at your Guidewire target so every mapping lines up with where the data must land.", YELLOW],
    ["D", "Explore the content", "Browse the tables and fields in your source data without needing a technical specialist.", GREEN],
    ["E", "Profile the quality", "See how complete and clean the real data is — before you commit to a mapping.", TEAL],
    ["F", "Work by client", "Each client engagement is kept separate, so teams can work in parallel with confidence.", ROSE],
])

# =====================================================================
# 9 - CAPABILITIES: MAP WITH AI (the core)
# =====================================================================
s = content_slide("What you can do · Step 3", "Map with AI — the heart of the tool", 9)
# left big statement panel
rect(s, 0.55, 2.0, 4.35, 4.35, INK2, rounded=True, radius=0.05)
text(s, 0.85, 2.35, 3.8, 0.5, "AI drafts.\nYou decide.", size=24, color=YELLOW, bold=True, face=HEAD, line_mult=1.0)
text(s, 0.85, 3.6, 3.8, 2.4,
     "The AI proposes the field-by-field mappings and transformation rules in seconds. Your experts stay in control — reviewing, adjusting and approving every one, so the result is trusted, not guessed.",
     size=12.5, color=DGRAY, line_mult=1.2)
# right stacked cards
items = [
    ["AI mapping suggestions", "Get a complete first-draft mapping for the fields you choose — instantly."],
    ["Review workspace", "Approve, reject or edit each mapping in one place; every change is tracked."],
    ["Code-value matching", "Legacy codes are matched to the right Guidewire values automatically."],
    ["Built-in checks", "The mappings are checked for gaps and low-confidence guesses before you move on."],
]
cx, cw2, cy, chh = 5.2, 7.6, 2.0, 1.0
for i, it in enumerate(items):
    y = cy + i*(chh + 0.12)
    rect(s, cx, y, cw2, chh, WHITE, rounded=True, line_color=LINE, line_w=1, radius=0.08)
    ellipse(s, cx + 0.25, y + chh/2 - 0.17, 0.34, [ORANGE, TANGERINE, GREEN, TEAL][i])
    text(s, cx + 0.75, y + 0.14, cw2 - 1.0, 0.35, it[0], size=13, color=INK, bold=True, face=HEAD)
    text(s, cx + 0.75, y + 0.5, cw2 - 1.0, 0.42, it[1], size=10.5, color=GRAY, line_mult=1.02)

# =====================================================================
# 10 - CAPABILITIES: BUILD, VALIDATE & DELIVER
# =====================================================================
s = content_slide("What you can do · Steps 4–6", "Build, validate and deliver with confidence", 10)
grid3x2(s, [
    ["A", "Generate the code", "Turn approved mappings into ready-to-run code that moves the data into the target.", ORANGE],
    ["B", "One-click deploy", "Push the code to the target database and track whether it succeeds.", TANGERINE],
    ["C", "Data-quality checks", "Check the loaded data for duplicates, missing values, bad codes and broken links.", YELLOW],
    ["D", "Quality dashboard", "See issues at a glance and drill into the exact records that need attention.", GREEN],
    ["E", "Audit trail & export", "Keep a full history of every change and export the mapping document to share.", TEAL],
    ["F", "Progress & usage", "Track project status and AI usage through clear, live dashboards.", ROSE],
])

# =====================================================================
# 11 - BUSINESS VALUE
# =====================================================================
s = content_slide("The value", "What this means for your business", 11)
grid3x2(s, [
    ["", "Faster delivery", "First-draft mappings in minutes turn weeks of manual work into days.", ORANGE],
    ["", "Lower cost", "Fewer specialist hours and far less rework across the program.", TANGERINE],
    ["", "Higher quality", "Data is validated against the target before go-live, not after.", YELLOW],
    ["", "You stay in control", "AI accelerates the work; your people approve every decision.", GREEN],
    ["", "Fully auditable", "Every mapping and change is tracked end-to-end for governance.", TEAL],
    ["", "Repeatable", "One consistent way of working across every client and source.", ROSE],
])

# =====================================================================
# 12 - TRUST, SECURITY & GOVERNANCE
# =====================================================================
s = content_slide("Trust", "Built for control and governance", 12)
grid2x2(s, [
    ["People stay in control", "AI suggests; your experts review and approve. Nothing is loaded on the AI's say-so alone.", ORANGE],
    ["A complete audit trail", "Who changed what, and when — captured across the whole project for full traceability.", TANGERINE],
    ["Each client kept separate", "Every engagement's data is strictly isolated, so nothing crosses between clients.", GREEN],
    ["Controlled access", "The tool is login-protected with administrator-managed users and permissions.", TEAL],
], circle=True)

# =====================================================================
# 13 - OUTCOMES
# =====================================================================
s = content_slide("The value", "The outcomes you can expect", 13)
text(s, 0.55, 1.72, 12.2, 0.55,
     "By making mapping faster and validation earlier, the Studio de-risks the part of the program that most often slips.",
     size=13.5, color=GRAY, line_mult=1.1)
outs = [
    ["Faster go-live", "Less time on mapping means the conversion stops being the bottleneck.", ORANGE],
    ["Fewer surprises", "Data issues are found and fixed early — not in the go-live weekend.", GREEN],
    ["Confident conversion", "A clear, approved, auditable trail from source to target.", TEAL],
]
cw, gx, gapX, gy, ch = 3.9, 0.55, 0.44, 2.55, 3.1
for i, o in enumerate(outs):
    x = gx + i*(cw+gapX)
    rect(s, x, gy, cw, ch, WHITE, rounded=True, line_color=LINE, line_w=1, radius=0.06)
    rect(s, x, gy, cw, 0.75, o[2], rounded=True, radius=0.09)
    text(s, x, gy, cw, 0.75, o[0], size=15, color=WHITE, bold=True, face=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.32, gy + 1.05, cw - 0.64, 1.8, o[1], size=13, color=GRAY, line_mult=1.2)

# =====================================================================
# 14 - ROADMAP
# =====================================================================
s = content_slide("What's next", "Where we are taking it", 14)
grid2x2(s, [
    ["Even faster on big sources", "Processing large data sources more quickly, so results come back sooner.", ORANGE],
    ["Handle bigger volumes", "Supporting ever-larger schemas and data sets with ease.", TANGERINE],
    ["Smarter suggestions", "Continual improvements to the quality of the AI's mapping proposals.", GREEN],
    ["Ongoing security hardening", "Continued investment to keep client data safe and access controlled.", TEAL],
], sand=True)

# =====================================================================
# 15 - CLOSING
# =====================================================================
s = slide(); bg(s, INK2)
rect(s, -1.5, 4.6, 5.5, 5.5, ORANGE, rounded=True, radius=0.08, rot=20)
rect(s, -0.6, 5.4, 4.0, 4.0, YELLOW, rounded=True, radius=0.08, rot=20)
brand_mark(s, 8.9, 0.75, 1.6)
text(s, 4.4, 2.35, 8.4, 0.7, "From legacy data to Guidewire —", size=30, color=WHITE, bold=True, face=HEAD)
text(s, 4.4, 3.0, 8.4, 0.7, "faster, cleaner, with confidence.", size=30, color=YELLOW, bold=True, face=HEAD)
text(s, 4.42, 3.95, 8.2, 1.0,
     "AI Data Conversion Studio brings AI-assisted mapping, automated build and early data validation into one guided workflow — with your team in control at every step.",
     size=14, color=DGRAY, line_mult=1.15)
text(s, 4.4, 5.5, 6, 0.6, "Thank you", size=22, color=WHITE, bold=True, face=HEAD)
footer(s, 15, True)

out = "AI_Data_Conversion_Studio_Overview.pptx"
prs.save(out)
print("WROTE", out, "slides:", len(prs.slides._sldIdLst))
