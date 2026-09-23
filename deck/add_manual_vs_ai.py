# -*- coding: utf-8 -*-
"""Insert a 'Manual vs. AI-assisted conversion' comparison slide into the Demo deck.

Matches the existing deck style: orange FD5108, red DC2626, card F5F7F8,
peach panel FFF5ED, Arial body / Georgia titles, footer 6B7280.
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml import parse_xml
import copy

SRC = "AI_Data_Conversion_Studio_Demo.pptx"

# ---- palette (from the existing deck) ----
ORANGE   = RGBColor(0xFD, 0x51, 0x08)
RED      = RGBColor(0xDC, 0x26, 0x26)
GREEN    = RGBColor(0x17, 0x8A, 0x4C)
INK      = RGBColor(0x00, 0x00, 0x00)
GRAY     = RGBColor(0x4A, 0x4A, 0x4A)
MUTE     = RGBColor(0x6B, 0x72, 0x80)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
CARD     = RGBColor(0xF5, 0xF7, 0xF8)   # manual (neutral) card
PEACH    = RGBColor(0xFF, 0xF5, 0xED)   # AI card / soft panel
LINE     = RGBColor(0xE4, 0xDE, 0xD9)
PEACHBDR = RGBColor(0xF3, 0xC9, 0xA6)

HEAD = "Georgia"
BODY = "Arial"

prs = Presentation(SRC)
W = prs.slide_width / 914400.0
H = prs.slide_height / 914400.0
DEFAULT = next(l for l in prs.slide_layouts if l.name == "DEFAULT")


def _empty_effect():
    return parse_xml('<a:effectLst xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"/>')


def rect(s, x, y, w, h, color, rounded=False, line_color=None, line_w=1.0, radius=0.08):
    st = MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE
    shp = s.shapes.add_shape(st, Inches(x), Inches(y), Inches(w), Inches(h))
    if rounded:
        try:
            shp.adjustments[0] = radius
        except Exception:
            pass
    if color is None:
        shp.fill.background()
    else:
        shp.fill.solid(); shp.fill.fore_color.rgb = color
    if line_color is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line_color; shp.line.width = Pt(line_w)
    shp._element.spPr.append(_empty_effect())
    return shp


def ellipse(s, x, y, d, color):
    shp = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    shp.fill.solid(); shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp._element.spPr.append(_empty_effect())
    return shp


def text(s, x, y, w, h, txt, size=12, color=INK, bold=False, face=BODY,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, spacing=None, line_mult=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    if line_mult:
        p.line_spacing = line_mult
    r = p.add_run(); r.text = txt
    r.font.size = Pt(size); r.font.color.rgb = color; r.font.bold = bold; r.font.name = face
    if spacing is not None:
        r._r.get_or_add_rPr().set('spc', str(int(spacing * 100)))
    return tb


def rich(s, x, y, w, h, segs, size=12, anchor=MSO_ANCHOR.MIDDLE, align=PP_ALIGN.LEFT):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = 0; tf.margin_right = 0; tf.margin_top = 0; tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align
    for (t, c, b) in segs:
        r = p.add_run(); r.text = t
        r.font.size = Pt(size); r.font.color.rgb = c; r.font.bold = b; r.font.name = BODY
    return tb


# ---------------------------------------------------------------- build slide
s = prs.slides.add_slide(DEFAULT)
s.background.fill.solid(); s.background.fill.fore_color.rgb = WHITE

# header
text(s, 0.55, 0.5, 11.5, 0.28, "WHY IT MATTERS", size=11, color=ORANGE, bold=True, face=BODY, spacing=2)
text(s, 0.55, 0.82, 11.9, 0.6, "Manual effort vs. AI-assisted conversion", size=27, color=INK, bold=True, face=HEAD)
text(s, 0.55, 1.47, 11.5, 0.45,
     "The same conversion work — the old, manual way versus with AI Data Conversion Studio.",
     size=12.5, color=GRAY, line_mult=1.1)

# column geometry
DIM_X, DIM_W = 0.55, 2.05
M_X,  M_W = 2.72, 4.95
A_X,  A_W = 7.86, 4.95

# header bars
HDR_Y, HDR_H = 2.12, 0.5
rect(s, M_X, HDR_Y, M_W, HDR_H, RED, rounded=True, radius=0.14)
text(s, M_X, HDR_Y, M_W, HDR_H, "Manual approach", size=13.5, color=WHITE, bold=True,
     face=BODY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
rect(s, A_X, HDR_Y, A_W, HDR_H, ORANGE, rounded=True, radius=0.14)
text(s, A_X, HDR_Y, A_W, HDR_H, "With AI Data Conversion Studio", size=13.5, color=WHITE, bold=True,
     face=BODY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# rows: (dimension, manual, ai)
rows = [
    ("Speed",
     "Weeks hand-mapping thousands of fields in spreadsheets.",
     "First-draft mappings generated in minutes."),
    ("Confidence",
     "Errors surface late — during testing, close to go-live.",
     "Fewer mapping issues; reviewers see the mapping up front."),
    ("Control",
     "Hard to review, easy to get wrong.",
     "Human-in-the-loop approves every mapping and rule."),
    ("Traceability",
     "Changes happen in silos — BA, config and DB drift apart.",
     "Full audit trail and logs at each decision step."),
    ("Consistency",
     "Varies by analyst and by engagement.",
     "One unified tool, repeatable across every client."),
]

ROW_Y0 = 2.78
ROW_H = 0.685
GAP = 0.09
CELL_H = ROW_H - GAP

for i, (dim, man, ai) in enumerate(rows):
    y = ROW_Y0 + i * ROW_H
    cy = y + CELL_H / 2
    # dimension label
    ellipse(s, DIM_X + 0.02, cy - 0.05, 0.1, ORANGE)
    text(s, DIM_X + 0.22, y, DIM_W - 0.22, CELL_H, dim, size=13, color=INK, bold=True,
         face=BODY, anchor=MSO_ANCHOR.MIDDLE)
    # manual cell
    rect(s, M_X, y, M_W, CELL_H, CARD, rounded=True, line_color=LINE, line_w=1, radius=0.10)
    ellipse(s, M_X + 0.22, cy - 0.15, 0.3, RED)
    text(s, M_X + 0.22, cy - 0.15, 0.3, 0.3, "✕", size=12, color=WHITE, bold=True,
         face=BODY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, M_X + 0.66, y, M_W - 0.86, CELL_H, man, size=10.5, color=GRAY,
         anchor=MSO_ANCHOR.MIDDLE, line_mult=1.04)
    # ai cell
    rect(s, A_X, y, A_W, CELL_H, PEACH, rounded=True, line_color=PEACHBDR, line_w=1, radius=0.10)
    ellipse(s, A_X + 0.22, cy - 0.15, 0.3, GREEN)
    text(s, A_X + 0.22, cy - 0.15, 0.3, 0.3, "✓", size=12, color=WHITE, bold=True,
         face=BODY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, A_X + 0.66, y, A_W - 0.86, CELL_H, ai, size=10.5, color=INK,
         anchor=MSO_ANCHOR.MIDDLE, line_mult=1.04)

# bottom takeaway bar
BB_Y = ROW_Y0 + len(rows) * ROW_H + 0.06
rect(s, 0.55, BB_Y, 12.23, 0.5, PEACH, rounded=True, line_color=PEACHBDR, line_w=1, radius=0.16)
rich(s, 0.85, BB_Y, 11.7, 0.5,
     [("The bottom line:  ", ORANGE, True),
      ("weeks of manual effort become minutes — with issues caught up front and every decision logged.",
       INK, False)],
     size=12, anchor=MSO_ANCHOR.MIDDLE)

# footer (match existing 8.5pt 6B7280)
text(s, 0.55, 7.046, 5.5, 0.28, "AI Data Conversion Studio", size=8.5, color=MUTE)

# ---------------------------------------------------------------- reorder: put new slide 4th (after Solution, before Demo)
sldIdLst = prs.slides._sldIdLst
ids = list(sldIdLst)          # order: [title, challenge, solution, demo, NEW]
new = ids[-1]
demo = ids[3]
sldIdLst.remove(new)
demo.addprevious(new)         # -> [title, challenge, solution, NEW, demo]

prs.save(SRC)
print("Saved. Slides now:", len(prs.slides._sldIdLst))
