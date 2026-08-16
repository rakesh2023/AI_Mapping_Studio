"""Unit tests for document_parser — one sample fixture per file type, plus the
data-profile output, dispatch, error handling, and the DB failure path."""
import io
import json

import pytest

from app.services import document_parser as P
from app.services.document_parser import DocumentParseError


# --------------------------------------------------------------------------- #
# Sample fixtures (built in-memory)
# --------------------------------------------------------------------------- #
def _csv_bytes():
    return b"claim_id,premium,state\n1,100.5,NY\n2,200,CA\n"

def _xlsx_bytes():
    from openpyxl import Workbook
    wb = Workbook(); ws = wb.active; ws.title = "Claims"
    ws.append(["ClaimId", "Premium", "State"]); ws.append([1, 100.5, "NY"]); ws.append([2, 200, "CA"])
    ws2 = wb.create_sheet("Policies"); ws2.append(["PolicyId", "Holder"]); ws2.append([10, "Acme Insurance"])
    buf = io.BytesIO(); wb.save(buf); return buf.getvalue()

def _json_records():
    return json.dumps([{"claim_id": 1, "premium": 100}, {"claim_id": 2, "premium": 250}]).encode()

def _json_nested():
    return json.dumps({"policy": {"number": 123, "coverages": ["A", "B"]}}).encode()

def _xml_bytes():
    return b"<policies><policy id='1'><number>123</number><premium>500</premium></policy></policies>"

def _sql_bytes():
    return b"CREATE TABLE Claim (ClaimId int, Premium decimal(18,2), State varchar(2));"

def _pdf_bytes(text="Policy Number 12345 Premium 500 Coverage A"):
    from reportlab.pdfgen import canvas
    buf = io.BytesIO(); c = canvas.Canvas(buf); c.drawString(72, 720, text); c.showPage(); c.save()
    return buf.getvalue()


# --------------------------------------------------------------------------- #
# Tabular parsers: DataFrame + data profile
# --------------------------------------------------------------------------- #
def test_parse_csv_returns_df_and_profile():
    r = P.parse_csv(_csv_bytes(), "claims.csv")
    assert r.kind == "structured" and r.dataframe is not None
    assert len(r.dataframe) == 2 and r.metadata["rowCount"] == 2 and r.metadata["columnCount"] == 3
    # profile has columns, dtypes, row count, and sample rows
    assert "Rows: 2" in r.profile and "claim_id" in r.profile and "Sample rows:" in r.profile
    assert r.tables and r.tables[0]["columns"][0]["name"] == "claim_id"
    assert "dtype" in r.tables[0]["columns"][0]


def test_parse_excel_profiles_each_sheet():
    r = P.parse_excel(_xlsx_bytes(), "book.xlsx")
    assert r.kind == "structured" and r.metadata["sheetCount"] == 2
    names = [t["name"] for t in r.tables]
    assert names == ["Claims", "Policies"]
    assert "Table: Claims" in r.text and "Table: Policies" in r.text
    assert r.dataframe is None            # multiple sheets -> no single df
    assert r.tables[0]["rowCount"] == 2


def test_parse_json_array_is_tabular():
    r = P.parse_json(_json_records(), "claims.json")
    assert r.kind == "structured" and r.dataframe is not None
    assert r.metadata["recordCount"] == 2 and "Rows: 2" in r.profile


def test_parse_json_nested_is_unstructured():
    r = P.parse_json(_json_nested(), "policy.json")
    assert r.kind == "unstructured" and r.metadata["topLevelType"] == "object"
    assert "coverages" in r.text and r.dataframe is None


def test_parse_sql_script_extracts_tables():
    r = P.parse_sql_script(_sql_bytes(), "schema.sql")
    assert r.kind == "structured" and r.metadata["tableCount"] == 1
    t = r.tables[0]
    assert t["name"].lower() == "claim"
    colnames = [c["name"] for c in t["columns"]]
    assert "ClaimId" in colnames and "Premium" in colnames
    assert "Table: Claim" in r.text


# --------------------------------------------------------------------------- #
# Unstructured parsers
# --------------------------------------------------------------------------- #
def test_parse_xml_flattens_text():
    r = P.parse_xml(_xml_bytes(), "policies.xml")
    assert r.kind == "unstructured" and r.metadata["rootTag"] == "policies"
    assert "123" in r.text and "500" in r.text


def test_parse_pdf_extracts_text():
    r = P.parse_pdf(_pdf_bytes(), "policy.pdf")
    assert r.kind == "unstructured" and r.metadata["pageCount"] == 1
    assert "12345" in r.text


# --------------------------------------------------------------------------- #
# Dispatch
# --------------------------------------------------------------------------- #
def test_dispatch_routes_by_extension():
    assert P.parse_document("x.csv", _csv_bytes()).kind == "structured"
    assert P.parse_document("x.json", _json_records()).kind == "structured"
    assert P.parse_document("x.xml", _xml_bytes()).kind == "unstructured"
    assert P.parse_document("x.pdf", _pdf_bytes()).kind == "unstructured"


def test_dispatch_unsupported_type_raises():
    with pytest.raises(DocumentParseError):
        P.parse_document("notes.txt", b"hello")


# --------------------------------------------------------------------------- #
# Graceful errors
# --------------------------------------------------------------------------- #
def test_empty_csv_raises():
    with pytest.raises(DocumentParseError):
        P.parse_csv(b"", "empty.csv")

def test_invalid_json_raises():
    with pytest.raises(DocumentParseError):
        P.parse_json(b"{ not json", "bad.json")

def test_invalid_xml_raises():
    with pytest.raises(DocumentParseError):
        P.parse_xml(b"<a><b></a>", "bad.xml")

def test_empty_sql_raises():
    with pytest.raises(DocumentParseError):
        P.parse_sql_script(b"   ", "empty.sql")


# --------------------------------------------------------------------------- #
# DB-aware orchestration: parse_and_store marks documents.status='failed'
# --------------------------------------------------------------------------- #
@pytest.fixture()
def _db(tmp_path, monkeypatch):
    monkeypatch.setenv("AIMS_DISABLE_DOTENV", "1")
    monkeypatch.setenv("AIMS_APP_DB", str(tmp_path / "parser_app.db"))
    from app.db.app_db import ensure_app_tables, connect
    ensure_app_tables()
    conn = connect()
    try:
        conn.execute("INSERT INTO users(email,password_hash,created_at) VALUES('p@x.com','h','t')")
        uid = conn.execute("SELECT id FROM users").fetchone()["id"]
        conn.execute("INSERT INTO clients(user_id,name,created_at) VALUES(?,?,?)", (uid, "C", "t"))
        cid = conn.execute("SELECT id FROM clients").fetchone()["id"]
        def mkdoc(name):
            cur = conn.execute(
                "INSERT INTO documents(user_id,client_id,filename,file_ext,status,created_at) "
                "VALUES(?,?,?,?, 'uploaded','t')", (uid, cid, name, name.rsplit(".", 1)[-1]))
            conn.commit()
            return cur.lastrowid
        yield {"connect": connect, "mkdoc": mkdoc}
    finally:
        conn.close()


def test_parse_and_store_marks_failed_on_error(_db):
    did = _db["mkdoc"]("bad.json")
    res = P.parse_and_store(did, "bad.json", b"{ not json")
    assert res is None
    conn = _db["connect"]()
    try:
        row = conn.execute("SELECT status, status_detail FROM documents WHERE id=?", (did,)).fetchone()
    finally:
        conn.close()
    assert row["status"] == "failed" and "Invalid JSON" in row["status_detail"]


def test_parse_and_store_success_leaves_status(_db):
    did = _db["mkdoc"]("claims.csv")
    res = P.parse_and_store(did, "claims.csv", _csv_bytes())
    assert res is not None and res.kind == "structured"
    conn = _db["connect"]()
    try:
        row = conn.execute("SELECT status FROM documents WHERE id=?", (did,)).fetchone()
    finally:
        conn.close()
    assert row["status"] == "uploaded"    # parser doesn't advance status on success
